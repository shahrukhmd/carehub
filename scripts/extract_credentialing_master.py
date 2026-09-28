"""Extract the Credentialing Master workbook into JSON for scripts/import-credentialing.ts.

Reads the "Credentialing Master" sheet (the source of truth) and uses each state tab only
for provider credentials and payer mailing addresses. The state tabs' status formulas are
offset by one row, so their status values are deliberately never read.

Usage: python scripts/extract_credentialing_master.py <workbook.xlsm> <out.json>
"""

import datetime
import json
import re
import sys

import openpyxl

MASTER_SHEET = "Credentialing Master"
FIRST_DATA_ROW = 4
HEADER_ROW = 3
FIRST_PROVIDER_COL = 8


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return {"t": "date", "v": v.date().isoformat()}
    if isinstance(v, datetime.time):
        return None
    text = str(v).strip()
    return {"t": "text", "v": text} if text else None


def main(path, out):
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    master = values[MASTER_SHEET]

    # Provider credential and display name per master column, from the state tabs' headers.
    credentials = {}
    addresses = {}
    for ws in formulas.worksheets:
        if ws.title == MASTER_SHEET:
            continue
        for col in range(FIRST_PROVIDER_COL, ws.max_column + 1):
            header, ref = ws.cell(1, col).value, ws.cell(2, col).value
            if not header or not isinstance(ref, str):
                continue
            m = re.match(r"='Credentialing Master'!([A-Z]+)\d+", ref)
            if not m:
                continue
            name, _, credential = str(header).partition("–")
            credentials[m.group(1)] = {"name": name.strip().rstrip(","), "credential": credential.strip() or None}
        for row in range(2, ws.max_row + 1):
            payer, address = ws.cell(row, 1).value, ws.cell(row, 4).value
            if payer and isinstance(address, str) and not address.startswith("="):
                addresses[(ws.title, str(payer).strip())] = address.strip()

    headers = {col: master.cell(HEADER_ROW, col).value for col in range(1, master.max_column + 1)}
    states = {}
    for row in range(FIRST_DATA_ROW, master.max_row + 1):
        state, payer = master.cell(row, 1).value, master.cell(row, 2).value
        if not state or not payer:
            continue
        state = str(state).strip()
        entry = states.setdefault(state, {"providers": {}, "lines": []})
        cells = {}
        for col in range(FIRST_PROVIDER_COL, master.max_column + 1):
            header = headers.get(col)
            if not header or not str(header).strip().endswith(f"({state})"):
                continue
            letter = openpyxl.utils.get_column_letter(col)
            if letter not in entry["providers"]:
                fallback = re.sub(r"\s*\(.*$", "", str(header)).strip()
                info = credentials.get(letter, {"name": fallback, "credential": None})
                entry["providers"][letter] = {"name": info["name"] or fallback, "credential": info["credential"]}
            cells[letter] = cell_value(master.cell(row, col).value)
        entry["lines"].append(
            {
                "payer": str(payer).strip(),
                "planType": str(master.cell(row, 3).value or "").strip(),
                "cpid": (str(master.cell(row, 4).value).strip() if master.cell(row, 4).value is not None else None),
                "address": addresses.get((state, str(payer).strip())),
                "group": cell_value(master.cell(row, 5).value),
                "edi": cell_value(master.cell(row, 6).value),
                "eft": cell_value(master.cell(row, 7).value),
                "cells": cells,
            }
        )

    with open(out, "w", encoding="utf-8") as fh:
        json.dump({"source": path, "states": states}, fh)
    total_cells = sum(len(l["cells"]) for s in states.values() for l in s["lines"])
    print(f"Extracted {len(states)} states, {sum(len(s['lines']) for s in states.values())} payer lines, "
          f"{sum(len(s['providers']) for s in states.values())} provider columns, {total_cells} status cells")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
